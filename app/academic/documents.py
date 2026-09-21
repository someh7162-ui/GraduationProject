"""Read-only parsers. Document text is data, never executable instructions."""
import io
import re
import zipfile
from xml.etree import ElementTree as ET
from fastapi import HTTPException
from .schemas import Course


def docx_pages(data):
    """Extract readable Word paragraphs and table rows without executing content."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if sum(entry.file_size for entry in entries) > 50 * 1024 * 1024:
            raise ValueError('Word 文档解压后过大')
        xml = archive.read('word/document.xml')

    root = ET.fromstring(xml)
    namespace = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    body = root.find('w:body', namespace)
    if body is None:
        return []

    blocks = []
    for child in body:
        if child.tag == f"{{{namespace['w']}}}p":
            text = ''.join(node.text or '' for node in child.findall('.//w:t', namespace)).strip()
            if text:
                blocks.append(text)
        elif child.tag == f"{{{namespace['w']}}}tbl":
            for row in child.findall('./w:tr', namespace):
                cells = []
                for cell in row.findall('./w:tc', namespace):
                    text = ''.join(node.text or '' for node in cell.findall('.//w:t', namespace)).strip()
                    if text:
                        cells.append(text)
                if cells:
                    blocks.append(' | '.join(cells))

    sections = []
    current = []
    heading = re.compile(r'^(?:[一二三四五六七八九十]+、|附件\s*\d+[:：])')
    for block in blocks:
        if heading.match(block) and current:
            sections.append('\n'.join(current))
            current = []
        current.append(block)
    if current:
        sections.append('\n'.join(current))
    return [(f'第 {index + 1} 节', text) for index, text in enumerate(sections)]


def text_pages(filename, data):
    suffix = filename.lower().rsplit('.', 1)[-1]
    if suffix == 'pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 100:
            raise ValueError('最多支持 100 页')
        return [(f'第 {i + 1} 页', page.extract_text() or '') for i, page in enumerate(reader.pages)]
    if suffix in ('txt', 'md'):
        return [('全文', data.decode('utf-8-sig'))]
    if suffix == 'docx':
        return docx_pages(data)
    rows = spreadsheet_rows(filename, data)
    return [(sheet, '\n'.join(' | '.join(str(v) for v in row) for row in values)) for sheet, values in rows]


def spreadsheet_rows(filename, data):
    if filename.lower().endswith('.xls'):
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        if sum(s.nrows * s.ncols for s in book.sheets()) > 250000:
            raise ValueError('表格过大')
        return [(s.name, [s.row_values(i) for i in range(s.nrows)]) for s in book.sheets()]
    if filename.lower().endswith('.xlsx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError('表格解压后过大')
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            result = []
            for s in book:
                if (s.max_row or 0) * (s.max_column or 0) > 250000:
                    raise ValueError('表格过大')
                result.append((s.title, [list(r) for r in s.iter_rows(values_only=True)]))
            return result
        finally:
            book.close()
    raise ValueError('支持 PDF、XLS、XLSX 文件')


def cell_name(index):
    result = ''
    while index:
        index, rem = divmod(index - 1, 26)
        result = chr(65 + rem) + result
    return result


def parse_transcript(filename, data, document_id):
    courses, facts, warnings = [], {}, []

    def source(locator):
        return {'document_id': document_id, 'filename': filename, 'locator': locator}

    pages = text_pages(filename, data)
    whole = '\n'.join(t for _, t in pages)
    years = re.search(r'(20\d{2})\s*[-—]\s*(20\d{2})', whole)
    year = f'{years[1]}-{years[2]}' if years else None
    for locator, text in pages:
        rank = re.search(r'专业班级名次[：:]\s*[（(]?\s*(\d+)\s*/\s*(\d+)', text)
        if rank:
            facts['unclassified_rank'] = {'value': {'rank': int(rank[1]), 'total': int(rank[2]), 'label': '专业班级名次'}, 'academic_year': year, 'sources': [source(locator)], 'confirmed': False}
        for key, pattern in [('college', r'院\s*系\s+(.+?)\s+专'), ('major', r'专\s*业\s+(.+?)\s+班'), ('class_name', r'班\s*级\s+(\S+)'), ('admission_year', r'入学时间\s+(\d{4})')]:
            match = re.search(pattern, text)
            if match:
                facts[key] = {'value': match[1].strip(), 'sources': [source(locator)], 'confirmed': False}
        for key, pattern in [('physical_education', r'体育成绩\s*\|\s*(\S+)'), ('moral_score', r'思想品德[（(]100分[）)]\s*\|\s*\|\s*(\d+)')]:
            match = re.search(pattern, text)
            if match:
                facts[key] = {'value': float(match[1]) if key == 'moral_score' else match[1], 'academic_year': year, 'sources': [source(locator)], 'confirmed': False}
    if filename.lower().endswith('.pdf'):
        pattern = re.compile(r'(20\d{2}-20\d{2})-([123])\s+(.+?)\s*(一体化课程|必修|限选|任选|公选)\s*\d+\s*[（(]学时[）)]\s*(\d+(?:\.\d+)?)\s+(\d+(?:\.\d+)?)')
        for locator, text in pages:
            for m in pattern.finditer(text):
                courses.append(Course(academic_year=m[1], term=int(m[2]), name=m[3].strip(), category=m[4], credits=float(m[5]), score=float(m[6]), sources=[source(locator)]).model_dump())
    else:
        for sheet, rows in spreadsheet_rows(filename, data):
            term = None
            for i, row in enumerate(rows):
                values = [str(v).strip() if v is not None else '' for v in row]
                if '第一学期' in values:
                    term = 1
                if '第二学期' in values:
                    term = 2
                if '课程' in values and i + 1 < len(rows) and term and year:
                    start = values.index('课程') + 1
                    for j in range(start, len(values)):
                        if not values[j] or j >= len(rows[i + 1]):
                            continue
                        value = rows[i + 1][j]
                        if isinstance(value, (int, float)):
                            courses.append(Course(academic_year=year, term=term, name=values[j], score=float(value), sources=[source(f'{sheet}!{cell_name(j + 1)}{i + 1}:{cell_name(j + 1)}{i + 2}')]).model_dump())
                # Standard vertical tables, e.g. exported XLSX grade reports.
                if '课程名称' in values and '成绩' in values:
                    cols = {v: j for j, v in enumerate(values) if v}
                    for k, record in enumerate(rows[i + 1:], i + 2):
                        def value(label, default=None):
                            col = cols.get(label)
                            return record[col] if col is not None and col < len(record) and record[col] not in ('', None) else default
                        if not value('课程名称') or not isinstance(value('成绩'), (int, float)):
                            continue
                        course_year = str(value('学年', year) or '')
                        course_term = value('学期', term)
                        if not course_year or not course_term:
                            warnings.append(f'{sheet}!{k} 行缺少学年或学期，请核对')
                            continue
                        courses.append(Course(academic_year=course_year, term=int(course_term), name=str(value('课程名称')), score=float(value('成绩')), credits=value('学分'), category=str(value('课程属性', '')), sources=[source(f'{sheet}!A{k}:{cell_name(len(record))}{k}')]).model_dump())
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    if str(value).strip() == '获奖情况':
                        award = next((str(v).strip() for v in row[j+1:] if v is not None and str(v).strip()), '')
                        if award:
                            facts['awards_text'] = {'value': award, 'academic_year': year, 'sources': [source(f'{sheet}!{cell_name(j + 1)}{i + 1}')], 'confirmed': False}
    if not courses:
        warnings.append('未自动识别课程；请核对格式或手动补录。扫描件 OCR 暂未支持。')
    if 'unclassified_rank' in facts:
        warnings.append('“专业班级名次”尚未明确排名类型与范围，不能直接用于综合测评排名核验。')
    return {'courses': courses, 'facts': facts, 'warnings': warnings, 'academic_year': year}


def course_key(course):
    return (course['academic_year'], course['term'], re.sub(r'\s+', '', course['name']).casefold())


def merge_courses(existing, incoming):
    import copy
    merged = {course_key(c): copy.deepcopy(c) for c in existing}
    conflicts = []
    for course in incoming:
        key = course_key(course)
        previous = merged.get(key)
        if previous is None:
            merged[key] = copy.deepcopy(course)
            continue
        incompatible = [field for field in ('score', 'credits', 'category') if previous.get(field) not in (None, '') and course.get(field) not in (None, '') and previous[field] != course[field]]
        if incompatible:
            conflicts.append({'course': course['name'], 'fields': incompatible, 'existing': previous, 'incoming': course})
            continue
        for field in ('credits', 'category'):
            if previous.get(field) in (None, ''):
                previous[field] = course.get(field)
        for source in course['sources']:
            if source not in previous['sources']:
                previous['sources'].append(source)
    if conflicts:
        raise HTTPException(409, {'message': '课程数据冲突，请核对后重新确认；原记录未覆盖', 'conflicts': conflicts})
    return list(merged.values())
